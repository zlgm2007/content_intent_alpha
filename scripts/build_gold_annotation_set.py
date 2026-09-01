"""生成「人工金标准标注集」Excel，用于拿到不受脏标签污染的可信准确率。

为什么需要它
------------
现状：所有意图指标（准确率 0.7196、二分 F1 0.8875）都是拿 ES 外部程序的打分当真值算的。
但已证实这套标签本身有系统性错误：

  - 低分档(0-2) 有 18.3% 含「求链接/多少钱/怎么买」这类明确购买信号却被判无意向
  - 高置信度(>0.9)「模型判 5 / 真值 0」的样本，抽样 12 条全是模型对、标签错
  - 1 分与 2 分在文本上找不到任何可辨识的系统差异

结论：**在脏标签上测出来的数字，既低估了模型，也无法指导下一步优化。**
必须有一批人工标注的金标准，才能回答「模型到底准不准」以及「该往哪改」。

抽样设计
--------
两个互补的层，各自回答不同的问题：

  P0 层（分层等量 252 条 = 每档 42 条）
      从留出集按真值分数分层、每层等量随机抽。
      作用：**无偏估计**各档准确率，并按真实分布加权还原整体准确率。
      为什么等量而不是按比例：0 分只占 4.7%，按比例抽只有 12 条，
      根本估不准它的准确率；而 0 分恰恰是标签问题最大的一档。

  P1 层（高置信度分歧样本 250 条）
      模型置信度 > 0.75 且与外部标签相差 >= 2 分。
      作用：**诊断**「到底是模型错还是标签错」——这些是信息量最大的样本。
      若这批里大多数是标签错，就坐实了标签质量问题。

产出
----
data/intent/labeled/gold_annotation.xlsx
  Sheet1 填写说明：打分标准（从数据归纳，比外部程序的标准更自洽）
  Sheet2 待标注  ：带下拉框，只需填一列

用法
----
    /opt/miniconda3/bin/python3 scripts/build_gold_annotation_set.py

依赖 /tmp/intent_logits_cache.npz（由 scripts/sweep_intent_prior_alpha.py 生成）。
若缓存不存在会自动重算 logits（约 4 分钟）。
"""

import csv
import json
import os
import random
import re

import numpy as np
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HOLDOUT = os.path.join("data", "intent", "labeled", "holdout.jsonl")
# sweep_intent_prior_alpha.py 的缓存，**只读**，不写回，避免覆盖它的 key 格式
# （它用 "路径|模型mtime|条数" 作键，写坏了会害它每次都重算）
CACHE = "/tmp/intent_logits_cache.npz"
OWN_CACHE = "/tmp/intent_logits_cache_gold.npz"
OUT = os.path.join("data", "intent", "labeled", "gold_annotation.xlsx")

P0_PER_CLASS = 42      # P0 层每档抽样数
P1_TOTAL = 250         # P1 层总数
P1_MIN_CONF = 0.75     # P1 层置信度门槛
P1_MIN_GAP = 2         # P1 层模型与标签最小分歧
SEED = 20260901

# 强购买意向信号（仅作参考列，提示标注者注意）
STRONG = re.compile(
    r"链接|多少钱|多钱|什么价|报价|怎么买|哪里买|哪买|在哪买|求推荐|有推荐"
    r"|推荐一下|推荐款|推荐一款|想要|想买|我要买|怎么卖|有货吗"
    r"|能给.*方式|联系(方式|一下)|私我"
)

# Excel 不接受部分控制字符
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(s: str, limit: int = None) -> str:
    s = ILLEGAL.sub("", str(s or "")).replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s{2,}", " ", s).strip()
    return s[:limit] if limit else s


def split_note_comment(text: str):
    """把组装文本拆回 (正文, 评论)。打分对象是评论，正文只作上下文。"""
    if "【评论】" in text:
        note, comment = text.split("【评论】", 1)
    else:
        note, comment = text, ""
    note = note.replace("【正文】", "").replace("【标题】", "").strip()
    return note, comment.strip()


def load_logits(rows):
    """取 logits (N,6)，优先复用 sweep 脚本的缓存。

    注意 IntentInference **没有** predict_proba 方法，只有 predict（返回 dict，
    丢掉了原始 logits）。要拿 logits 必须走 tokenizer + runner.run，与
    sweep_intent_prior_alpha.py 保持一致。
    """
    for path in (CACHE, OWN_CACHE):
        if os.path.exists(path):
            z = np.load(path, allow_pickle=True)
            lg = z["logits"]
            if lg.shape[0] == len(rows):
                print(f"  复用缓存 {lg.shape} <- {path}")
                return lg
            print(f"  {path} 条数 {lg.shape[0]} != {len(rows)}，跳过")

    from src.intent.inference import IntentInference  # 延迟导入，避免无谓开销

    texts = [r["text"] for r in rows]
    print(f"  计算 logits（{len(texts)} 条，约需数分钟）…", flush=True)
    inf = IntentInference("configs/intent.yaml")
    batch = 64
    outs = []
    for i in range(0, len(texts), batch):
        enc = inf.tokenizer(
            texts[i : i + batch],
            max_length=inf.max_length,
            padding="max_length",
            truncation=True,
            return_tensors="np",
        )
        outs.append(np.asarray(inf.runner.run(enc["input_ids"], enc["attention_mask"]),
                               dtype=np.float64))
        if (i // batch) % 20 == 0:
            print(f"    推理 {i}/{len(texts)}", flush=True)
    lg = np.vstack(outs)
    np.savez(OWN_CACHE, key="gold", logits=lg)
    print(f"  已缓存 -> {OWN_CACHE}")
    return lg


def apply_prior_correction(lg: np.ndarray) -> np.ndarray:
    """把推理端的先验修正（logits += alpha * bias）应用到离线 logits。

    为什么必须加这一步：线上的 IntentInference.predict 会做
        logits += prior_alpha * prior_bias
    而本脚本直接对**原始** logits 取 argmax，得到的是未修正的预测。
    configs/intent.yaml 的 prior_correction_alpha 一旦 >0（2026-09-01 起设为 0.8），
    表里「模型预测」「模型置信度」两列就会与线上实际行为不一致。

    后果很实际：eval_on_gold.py 评测时**读表内的静态列**、不重新推理，
    所以报出的「模型 vs 人工准确率」会停留在旧 alpha 的水平，与线上差 1.28pp。
    人工标 500 条成本高，评测口径必须先对齐再标，否则整批数据白标。
    """
    cfg = yaml.safe_load(open("configs/intent.yaml", encoding="utf-8"))
    inf = cfg.get("inference") or {}
    alpha = float(inf.get("prior_correction_alpha") or 0.0)
    if alpha <= 0:
        print(f"  先验修正关闭（alpha={alpha}），使用原始 logits")
        return lg
    with open(inf["prior_file"], "r", encoding="utf-8") as f:
        bias = np.array(json.load(f)["bias"], dtype=np.float64)
    print(f"  应用先验修正 alpha={alpha}  bias={np.round(bias, 3).tolist()}")
    return lg + alpha * bias


def main():
    print("读取留出集…")
    rows = [json.loads(l) for l in open(HOLDOUT, encoding="utf-8")]
    y = np.array([r["label"] for r in rows])
    n = len(rows)
    print(f"  {n} 条，分布 {dict(zip(*np.unique(y, return_counts=True)))}")

    print("准备 logits…")
    lg = apply_prior_correction(load_logits(rows))
    pred = lg.argmax(1)
    p = np.exp(lg - lg.max(1, keepdims=True))
    p /= p.sum(1, keepdims=True)
    conf = p.max(1)

    random.seed(SEED)
    chosen = {}  # idx -> 优先级

    # ---- P0 层：分层等量，无偏估计 ----
    for c in range(6):
        pool = [i for i in range(n) if y[i] == c]
        random.shuffle(pool)
        for i in pool[:P0_PER_CLASS]:
            chosen[i] = "P0"
    print(f"  P0 层（分层等量）{sum(1 for v in chosen.values() if v == 'P0')} 条")

    # ---- P1 层：高置信度分歧样本，用于诊断 ----
    cand = [
        i for i in range(n)
        if i not in chosen
        and conf[i] > P1_MIN_CONF
        and abs(pred[i] - y[i]) >= P1_MIN_GAP
    ]
    # 按真值档轮流取，保证每个档都有代表，而不是被多数类霸占
    by_class = {c: [i for i in cand if y[i] == c] for c in range(6)}
    for lst in by_class.values():
        lst.sort(key=lambda i: -conf[i])
    order, k = [], 0
    while len(order) < P1_TOTAL and any(by_class.values()):
        c = k % 6
        if by_class[c]:
            order.append(by_class[c].pop(0))
        k += 1
    for i in order:
        chosen[i] = "P1"
    print(f"  P1 层（高置信分歧）{len(order)} 条")

    # P0 在前，方便时间紧时只标 P0
    order_idx = sorted(chosen.keys(), key=lambda i: (chosen[i] != "P0", y[i], i))

    # ---- 写 Excel ----
    wb = Workbook()

    # Sheet1 填写说明
    ws = wb.active
    ws.title = "填写说明"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 104

    def put(a, b="", bold=False, fill=None):
        r = ws.max_row + 1 if ws.max_row > 1 or ws["A1"].value else 1
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        if bold:
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=2).font = Font(bold=True)
        if fill:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=fill)
        return r

    put("为什么要标这批数据", "", bold=True, fill="FAEEDA")
    put(
        "背景",
        "目前所有意图指标都是拿 ES 外部程序的打分当「标准答案」算的。但已证实这套答案本身有系统性错误："
        "低分档(0-2) 里有 18.3% 含「求链接/多少钱/怎么买」这类明确购买信号却被判为无意向；"
        "模型高置信度判 5 分而标签是 0 分的样本，抽样 12 条全部是模型对、标签错。"
        "所以在脏答案上测出的数字既低估了模型，也没法指导下一步优化。",
    )
    put(
        "这批数据的用途",
        "标完之后可以回答两个之前回答不了的问题：(1) 模型真实的准确率是多少；"
        "(2) 模型判错的那些，到底是模型不行还是标签不行。",
    )
    put("要标多少", f"共 {len(order_idx)} 条。P0 层 {P0_PER_CLASS * 6} 条必标（用来算可信准确率），"
                  f"P1 层 {len(order)} 条建议标（用来诊断分歧）。时间紧就只标 P0 层。")
    put("预计耗时", "P0 层约 40-60 分钟，全部约 1.5-2 小时。")

    put("")
    put("打分标准（0-5 分）", "", bold=True, fill="FAEEDA")
    put("看什么", "只看【评论原文】这一列。正文只是上下文，用来理解评论在说什么。")
    put(
        "0 分",
        "完全无关。没提到任何商品/品牌/品类，纯闲聊、灌水、与购买毫无关系。"
        "例：「在达州是不是还有一套房子」「没找到呢」",
    )
    put(
        "1 分",
        "提到了品牌或品类，但没有任何想要、询问的表示。"
        "例：「用惯哪个选哪个，都那么回事」「大品牌值得信赖」",
    )
    put(
        "2 分",
        "提到具体商品并流露出兴趣，但没有明确购买意图。"
        "例：「蹲」「不错可以买」（泛泛而言）「真的吗，这么便宜」",
    )
    put(
        "3 分",
        "有明确咨询意向：问价格、问推荐、问好不好用、问怎么选。"
        "例：「什么牌子介绍一下」「5pro 和栗峰 sl 怎么选」「这个推荐吗」",
    )
    put(
        "4 分",
        "有明确购买意愿：说想买、问在哪买、问有没有优惠、问有没有货。"
        "例：「想买个冰箱或冰柜，有推荐的吗」「大学刚毕业想买个 1 到 2 万的」",
    )
    put(
        "5 分",
        "已到成交环节：求链接、要联系方式、说要下单、问怎么付款。"
        "例：「求链接」「主播裤子链接能给一下吗」「方便留个联系方式吗」",
    )
    put(
        "拿不准时",
        "按「这个人离下单还有多远」判断，别纠结措辞。"
        "边界样本在备注列写一句理由，比硬凑一个分更有价值。",
    )

    put("")
    put("表格里各列怎么用", "", bold=True, fill="FAEEDA")
    put("正文摘要", "只作上下文，帮助你理解评论的语境。")
    put("评论原文", "★ 打分对象，看这一列就够了。")
    put("外部标签", "ES 外部程序给的分数，仅供参考——它有可能是错的，别被它带偏。")
    put("模型预测", "当前模型的预测，同样仅供参考。")
    put("模型置信度", "0-1，越高表示模型越确定。")
    put("强信号", "评论里是否出现「求链接/多少钱/怎么买」这类显性购买表达，命中显示 Y。")
    put("你的打分", "★ 唯一需要你填的列，填 0-5 的整数。已加下拉框。")
    put("备注", "选填。拿不准、或觉得这条本身没法打分时写一句。")

    put("")
    put("一个提醒", "", bold=True, fill="FCEBEB")
    put(
        "别被「外部标签」带偏",
        "标之前先别看它。如果先看了再标，会不自觉地向它靠拢，"
        "这批数据就失去了「独立标准答案」的意义，白标了。",
    )

    # Sheet2 待标注
    ws2 = wb.create_sheet("待标注")
    headers = [
        "序号", "优先级", "正文摘要", "评论原文（★打分看这列）",
        "外部标签", "模型预测", "模型置信度", "强信号",
        "★你的打分(0-5)", "备注",
    ]
    ws2.append(headers)
    for col, w in zip("ABCDEFGHIJ", [6, 8, 46, 54, 10, 10, 11, 8, 15, 26]):
        ws2.column_dimensions[col].width = w

    head_fill = PatternFill("solid", fgColor="E6F1FB")
    for c in range(1, len(headers) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws2.row_dimensions[1].height = 30

    for seq, i in enumerate(order_idx, 1):
        note, comment = split_note_comment(rows[i]["text"])
        ws2.append([
            seq,
            chosen[i],
            clean(note, 90),
            clean(comment, 200),
            int(y[i]),
            int(pred[i]),
            round(float(conf[i]), 3),
            "Y" if STRONG.search(comment) else "",
            None,
            None,
        ])

    for r in range(2, ws2.max_row + 1):
        for c in (3, 4, 10):
            ws2.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
        for c in (1, 2, 5, 6, 7, 8, 9):
            ws2.cell(row=r, column=c).alignment = Alignment(
                vertical="top", horizontal="center"
            )
        ws2.row_dimensions[r].height = 42

    dv = DataValidation(
        type="list", formula1='"0,1,2,3,4,5"', allow_blank=True, showDropDown=False
    )
    dv.error = "请填 0 到 5 的整数"
    dv.errorTitle = "超出范围"
    ws2.add_data_validation(dv)
    dv.add(f"I2:I{ws2.max_row}")

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    n_p0 = sum(1 for v in chosen.values() if v == "P0")
    n_p1 = sum(1 for v in chosen.values() if v == "P1")
    print(f"\n已生成 {OUT}")
    print(f"  P0 层（分层等量，必标）  {n_p0} 条")
    print(f"  P1 层（高置信分歧）      {n_p1} 条")
    print(f"  合计                    {n_p0 + n_p1} 条")

    # 顺带导出一份 CSV 备份，方便用其他工具打开
    csv_path = OUT.replace(".xlsx", ".csv")
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for seq, i in enumerate(order_idx, 1):
            note, comment = split_note_comment(rows[i]["text"])
            w.writerow([
                seq, chosen[i], clean(note, 90), clean(comment, 200),
                int(y[i]), int(pred[i]), round(float(conf[i]), 3),
                "Y" if STRONG.search(comment) else "", "", "",
            ])
    print(f"  CSV 备份                {csv_path}")


if __name__ == "__main__":
    main()
